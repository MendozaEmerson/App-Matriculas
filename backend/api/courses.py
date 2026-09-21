import pandas as pd
from io import BytesIO
from fastapi import APIRouter, File, UploadFile, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from core.database import get_db
from models.course import Course, CourseGroup
import io
import base64
from pydantic import BaseModel

router = APIRouter()

class UploadBase64Request(BaseModel):
    filename: str
    base64_data: str

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

@router.get("/template")
def download_courses_template():
    """Genera y descarga una plantilla Excel para subir cursos"""
    df = pd.DataFrame(columns=[
        'Código Teórico', 'Código Laboratorio', 'Nombre', 'Año', 
        'Semestre', 'Grupo', 'Horario', 'Vacantes'
    ])
    
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Cursos')
        worksheet = writer.sheets['Cursos']
        
        # Ajustar anchos y estilizar la cabecera (Tabla azul)
        for i, col in enumerate(df.columns, 1):
            col_letter = chr(64 + i)
            worksheet.column_dimensions[col_letter].width = len(col) + 8
            
            cell = worksheet.cell(row=1, column=i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") # Azul Tailwind
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    stream.seek(0)
    
    return StreamingResponse(
        stream, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_cursos.xlsx"}
    )

@router.post("/upload_b64")
async def upload_courses_b64(request: UploadBase64Request, db: Session = Depends(get_db)):
    if not request.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos Excel (.xlsx)")
        
    try:
        contents = base64.b64decode(request.base64_data)
        return _process_courses_excel(contents, request.filename, db)
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error procesando el Excel: {str(e)}")

@router.post("/upload")
async def upload_courses(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Recibe un archivo Excel (.xlsx).
    Agrupa los cursos y sus grupos (A, B, C...) insertándolos en la base de datos local SQLite.
    """
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos Excel (.xlsx)")
        
    contents = await file.read()
    try:
        return _process_courses_excel(contents, file.filename, db)
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error procesando el Excel: {str(e)}")

def _process_courses_excel(contents: bytes, filename: str, db: Session):
        # Leemos el Excel en memoria con Pandas
        df = pd.read_excel(BytesIO(contents))
        
        courses_created = 0
        groups_created = 0
        
        for index, row in df.iterrows():
            # Buscar el código teórico con múltiples posibles nombres
            theoretical_code = str(row.get('Código Teórico', row.get('codigo_teorico', row.get('Codigo Curso Teorico', '')))).strip()
            
            if not theoretical_code or theoretical_code == 'nan':
                continue
                
            # Validamos si el curso ya existe en la DB
            course = db.query(Course).filter(Course.theoretical_code == theoretical_code).first()
            if not course:
                year_val = row.get('Año', row.get('year', 1))
                semester_val = row.get('Semestre', row.get('semester', 1))
                lab_code = str(row.get('Código Laboratorio', row.get('codigo_laboratorio', row.get('Codigo Curso Laboratorio', '')))).strip()
                name = str(row.get('Nombre', row.get('name', row.get('Nombre del Curso', '')))).strip()
                
                course = Course(
                    theoretical_code=theoretical_code,
                    lab_code=lab_code if lab_code != 'nan' else '',
                    name=name if name != 'nan' else '',
                    year=int(year_val) if not pd.isna(year_val) else 1,
                    semester=int(semester_val) if not pd.isna(semester_val) else 1
                )
                db.add(course)
                db.flush()
                courses_created += 1
                
            # Extraer grupos, horarios y vacantes (pueden venir separados por comas)
            groups_raw = str(row.get('Grupo', row.get('Grupos', row.get('group', 'A')))).strip()
            schedules_raw = str(row.get('Horario', row.get('Horario por Grupo', row.get('horario', '')))).strip()
            vacancies_raw = str(row.get('Vacantes', row.get('Vacantes por Grupo', row.get('vacancies', '0')))).strip()
            
            # Dividir por comas en caso el usuario ponga varios grupos en una fila (Ej: "A,B,C")
            groups_list = [g.strip() for g in groups_raw.split(',')] if groups_raw and groups_raw != 'nan' else []
            schedules_list = [s.strip() for s in schedules_raw.split(',')] if schedules_raw and schedules_raw != 'nan' else []
            vacancies_list = [v.strip() for v in vacancies_raw.split(',')] if vacancies_raw and vacancies_raw != 'nan' else []
            
            # Iterar sobre cada grupo y crearlo
            for i, group_name in enumerate(groups_list):
                if not group_name: continue
                
                schedule = schedules_list[i] if i < len(schedules_list) else ''
                vacancies_str = vacancies_list[i] if i < len(vacancies_list) else '0'
                
                try:
                    vacancies_num = int(vacancies_str)
                except ValueError:
                    vacancies_num = 0

                existing_group = db.query(CourseGroup).filter(
                    CourseGroup.course_id == course.id,
                    CourseGroup.name == group_name
                ).first()
                
                if not existing_group:
                    group = CourseGroup(
                        course_id=course.id,
                        name=group_name,
                        schedule_range=schedule,
                        initial_vacancies=vacancies_num,
                        available_vacancies=vacancies_num
                    )
                    db.add(group)
                    groups_created += 1
                
        # Confirmamos todos los cambios en la base de datos
        db.commit()
        
        return {
            "filename": filename, 
            "status": "success", 
            "message": f"Procesamiento exitoso: Se guardaron {courses_created} cursos y {groups_created} grupos en la base de datos."
        }
