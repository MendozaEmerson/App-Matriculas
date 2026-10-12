import pandas as pd
from io import BytesIO
import unicodedata
from fastapi import APIRouter, File, UploadFile, Depends, HTTPException
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from core.database import get_db
from models.student import Student
from models.upload_log import UploadLog
import io
import base64
from pydantic import BaseModel

router = APIRouter()

class UploadBase64Request(BaseModel):
    filename: str
    base64_data: str

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment

@router.get("/template")
def download_students_template():
    """Genera y descarga una plantilla Excel para subir estudiantes"""
    df = pd.DataFrame(columns=[
        'CUI', 'DNI', 'Nombres', 'Apellido Paterno', 'Apellido Materno', 
        'Correo', 'Codigos cursos inscritos'
    ])
    
    stream = io.BytesIO()
    with pd.ExcelWriter(stream, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Estudiantes')
        worksheet = writer.sheets['Estudiantes']
        
        # Ajustar anchos y estilizar la cabecera (Tabla verde)
        for i, col in enumerate(df.columns, 1):
            col_letter = chr(64 + i)
            worksheet.column_dimensions[col_letter].width = len(col) + 12
            
            cell = worksheet.cell(row=1, column=i)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="059669", end_color="059669", fill_type="solid") # Verde Tailwind
            cell.alignment = Alignment(horizontal="center", vertical="center")
            
    stream.seek(0)
    
    return StreamingResponse(
        stream, 
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla_estudiantes.xlsx"}
    )

def clean_string(text):
    """Limpia tildes y pasa a minúsculas para generar correos automáticamente"""
    if pd.isna(text): return ""
    text = str(text).strip().lower()
    return ''.join(c for c in unicodedata.normalize('NFD', text) if unicodedata.category(c) != 'Mn')

@router.post("/upload_b64")
async def upload_students_b64(request: UploadBase64Request, db: Session = Depends(get_db)):
    if not request.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos Excel (.xlsx)")
        
    try:
        contents = base64.b64decode(request.base64_data)
        return _process_students_excel(contents, request.filename, db)
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error procesando el Excel: {str(e)}")

@router.post("/upload")
async def upload_students(file: UploadFile = File(...), db: Session = Depends(get_db)):
    """
    Recibe un archivo Excel (.xlsx) de estudiantes.
    Si el correo no existe, lo genera automáticamente basado en nombre y apellido.
    """
    if not file.filename.endswith('.xlsx'):
        raise HTTPException(status_code=400, detail="Solo se aceptan archivos Excel (.xlsx)")
        
    contents = await file.read()
    try:
        return _process_students_excel(contents, file.filename, db)
    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Error procesando el Excel: {str(e)}")

def _process_students_excel(contents: bytes, filename: str, db: Session):
    df = pd.read_excel(BytesIO(contents))
    
    total_rows = len(df)
    students_created = 0
    duplicated_students_count = 0
    duplicated_items = []
    
    for index, row in df.iterrows():
        cui = str(row.get('CUI', row.get('cui', ''))).strip()
        
        # Validamos que sea de 8 dígitos según lo indicado
        if not cui or cui == 'nan' or len(cui) != 8:
            continue
            
        student = db.query(Student).filter(Student.cui == cui).first()
        
        first_name = str(row.get('Nombres', row.get('nombres', ''))).strip()
        paternal = str(row.get('Apellido Paterno', row.get('apellido_paterno', ''))).strip()
        maternal = str(row.get('Apellido Materno', row.get('apellido_materno', ''))).strip()
        full_name = f"{paternal} {maternal} {first_name}".strip() if paternal or first_name else f"Estudiante {cui}"
        
        if student:
            duplicated_students_count += 1
            duplicated_items.append({
                "identifier": cui,
                "name": f"{student.paternal_surname} {student.first_name}",
                "reason": "Estudiante con CUI existente en el padrón"
            })
        else:
            # Generador automático de correos (Ej: jperez@unsa.edu.pe)
            email = str(row.get('Correo', row.get('correo', ''))).strip()
            if not email or email == 'nan':
                first_letter = clean_string(first_name)[0] if first_name else ''
                clean_pat = clean_string(paternal)
                email = f"{first_letter}{clean_pat}@unsa.edu.pe"
            
            # Asumimos que los cursos vienen separados por comas en la columna "Codigos cursos inscritos"
            courses_raw = str(row.get('Codigos cursos inscritos', row.get('cursos_inscritos', ''))).strip()
            if courses_raw and courses_raw != 'nan':
                enrolled_courses = [c.strip() for c in courses_raw.split(',')]
            else:
                enrolled_courses = []
            
            student = Student(
                cui=cui,
                dni=str(row.get('DNI', row.get('dni', ''))).strip(),
                first_name=first_name,
                paternal_surname=paternal,
                maternal_surname=maternal,
                institutional_email=email,
                enrolled_course_codes=enrolled_courses
            )
            db.add(student)
            students_created += 1
            
    # Guardar registro en upload_logs
    upload_log = UploadLog(
        upload_type='students',
        filename=filename,
        total_rows=total_rows,
        created_count=students_created,
        duplicated_count=duplicated_students_count,
        duplicated_items=duplicated_items
    )
    db.add(upload_log)
    db.commit()
    
    return {
        "filename": filename,
        "status": "success",
        "total_rows": total_rows,
        "created_count": students_created,
        "duplicated_count": duplicated_students_count,
        "duplicated_items": duplicated_items,
        "message": f"Procesamiento completado: Se registraron {students_created} estudiantes nuevos y se omitieron {duplicated_students_count} duplicados."
    }

@router.get("/upload_history")
def get_students_upload_history(db: Session = Depends(get_db)):
    """Obtiene el historial de subidas de estudiantes y auditoría de duplicados"""
    logs = db.query(UploadLog).filter(UploadLog.upload_type == 'students').order_by(UploadLog.created_at.desc()).limit(15).all()
    return [
        {
            "id": log.id,
            "upload_type": log.upload_type,
            "filename": log.filename,
            "total_rows": log.total_rows,
            "created_count": log.created_count,
            "duplicated_count": log.duplicated_count,
            "duplicated_items": log.duplicated_items or [],
            "created_at": log.created_at.isoformat() if log.created_at else ""
        }
        for log in logs
    ]
